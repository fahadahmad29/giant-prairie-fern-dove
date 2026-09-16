#!/usr/bin/env python3
"""TMU-format project report: Multiple Disease Prediction System — Shorya Agarwal."""
from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor, Emu, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement
from docx.enum.style import WD_STYLE_TYPE

FIG = Path("/workspace/artifacts/report_figures")
OUT = Path("/workspace/artifacts/Multiple_Disease_Prediction_System_Project_Report_Shorya_Agarwal.docx")

NAVY = RGBColor(0x1B, 0x36, 0x5D)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
GOLD = RGBColor(0xB8, 0x94, 0x4A)
BLACK = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x47, 0x55, 0x69)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY_HEX = "1B365D"
TEAL_HEX = "0F766E"
GOLD_HEX = "B8944A"
LIGHT_HEX = "F1F5F9"
SOFT_HEX = "E8EEF5"

STUDENT = "Shorya Agarwal"
ENROLL = "TCA2568564"
PROGRAM = "Bachelors in Computer Application"
SECTION = "G"
SEMESTER = "5th"
FACULTY = "Faculty of Engineering & Computing Sciences"
UNIV = "Teerthanker Mahaveer University, Moradabad"
TITLE_LONG = "Development of a Multiple Disease Prediction System Using Machine Learning"
TITLE_SHORT = "Multiple Disease Prediction System"
GUIDE1 = "Mr. Savood Ahmad"
GUIDE2 = "Ms. Sneha"
DATE_STR = "17-09-2026"
PLACE = "Moradabad"
MONTH = "September, 2026"


def set_run_font(run, name="Times New Roman", size=12, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def set_paragraph_spacing(p, before=0, after=8, line=1.5, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    return p


def add_text(doc, text, *, size=12, bold=False, italic=False, align="justify",
             before=0, after=8, line=1.5, font="Times New Roman", color=None,
             space_first=0, keep_with_next=False):
    p = doc.add_paragraph()
    al = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    set_paragraph_spacing(p, before, after, line, al)
    if space_first:
        p.paragraph_format.first_line_indent = Cm(space_first)
    p.paragraph_format.keep_with_next = keep_with_next
    run = p.add_run(text)
    set_run_font(run, font, size, bold, italic, color)
    return p


def add_mixed(doc, parts, *, align="justify", before=0, after=8, line=1.5, size=12):
    """parts: list of (text, bold, italic, color or None)."""
    p = doc.add_paragraph()
    al = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    set_paragraph_spacing(p, before, after, line, al)
    for item in parts:
        text, bold, italic = item[0], item[1], item[2]
        color = item[3] if len(item) > 3 else None
        sz = item[4] if len(item) > 4 else size
        run = p.add_run(text)
        set_run_font(run, "Times New Roman", sz, bold, italic, color)
    return p


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell, **kwargs):
    """kwargs: top/left/bottom/right = {sz, color, val}"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        spec = kwargs.get(edge, {"sz": "4", "color": "CBD5E1", "val": "single"})
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), spec.get("val", "single"))
        el.set(qn("w:sz"), spec.get("sz", "4"))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), spec.get("color", "CBD5E1"))
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def cell_text(cell, text, *, bold=False, size=11, align="left", color=None,
              fill=None, font="Times New Roman", italic=False, center_v=True):
    cell.text = ""
    p = cell.paragraphs[0]
    al = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    p.alignment = al
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, font, size, bold, italic, color or BLACK)
    if fill:
        shade_cell(cell, fill)
    if center_v:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)
    return p


def style_table(table, header=True, header_fill=NAVY_HEX, col_widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "CBD5E1")
        borders.append(el)
    tblPr.append(borders)
    # prevent row split
    for row in table.rows:
        tr = row._tr
        trPr = tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
    if header and table.rows:
        for cell in table.rows[0].cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.color.rgb = WHITE
                    run.bold = True
            shade_cell(cell, header_fill)
    if col_widths:
        table.autofit = False
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)


def add_hline(paragraph, color=NAVY_HEX, sz="12"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def set_page_number_format(section, fmt="decimal", start=None):
    sectPr = section._sectPr
    # remove existing pgNumType
    for child in list(sectPr):
        if child.tag == qn("w:pgNumType"):
            sectPr.remove(child)
    pgNumType = OxmlElement("w:pgNumType")
    pgNumType.set(qn("w:fmt"), fmt)
    if start is not None:
        pgNumType.set(qn("w:start"), str(start))
    sectPr.append(pgNumType)


def add_page_field(paragraph, instr="PAGE"):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = instr
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instrText)
    run._r.append(fld2)
    set_run_font(run, "Times New Roman", 10, False, False, MUTED)


def setup_header_footer(section, show_header=True, roman=False):
    header = section.header
    footer = section.footer
    header.is_linked_to_previous = False
    footer.is_linked_to_previous = False
    header.paragraphs[0].clear()
    footer.paragraphs[0].clear()
    if show_header:
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(hp, 0, 2, 1.0)
        r = hp.add_run(f"{TITLE_SHORT}  •  Teerthanker Mahaveer University")
        set_run_font(r, "Times New Roman", 9, False, True, NAVY)
        add_hline(hp, GOLD_HEX, "8")
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(fp, 4, 0, 1.0)
    add_hline(fp, GOLD_HEX, "8")
    r1 = fp.add_run("Page ")
    set_run_font(r1, "Times New Roman", 10, False, False, MUTED)
    add_page_field(fp, "PAGE")
    r2 = fp.add_run("  |  BCA Project Report  |  ")
    set_run_font(r2, "Times New Roman", 10, False, False, MUTED)
    r3 = fp.add_run(f"{STUDENT} ({ENROLL})")
    set_run_font(r3, "Times New Roman", 10, False, False, MUTED)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(8)

    for level, size in ((1, 16), (2, 14), (3, 12)):
        st = doc.styles[f"Heading {level}"]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = NAVY
        st.font.italic = False
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        st._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        sp = st.paragraph_format
        sp.space_before = Pt(16 if level == 1 else 12)
        sp.space_after = Pt(8)
        sp.keep_with_next = True
        sp.line_spacing = 1.15


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p


def bullet(doc, text, *, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(1.0)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        set_run_font(r, "Times New Roman", 12, True, False, BLACK)
        r2 = p.add_run(text)
        set_run_font(r2, "Times New Roman", 12, False, False, BLACK)
    else:
        r = p.add_run(text)
        set_run_font(r, "Times New Roman", 12, False, False, BLACK)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(1.0)
    r = p.add_run(text)
    set_run_font(r, "Times New Roman", 12)
    return p


def caption(doc, text):
    p = add_text(doc, text, size=11, italic=True, align="center", before=4, after=14, line=1.15, color=NAVY)
    return p


def add_figure(doc, filename, width, cap):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, 10, 4, 1.0)
    run = p.add_run()
    run.add_picture(str(FIG / filename), width=Inches(width))
    caption(doc, cap)


def page_break(doc):
    doc.add_page_break()


def prevent_table_row_split(table):
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)


# ---------------------------------------------------------------------------
# Cover and front matter
# ---------------------------------------------------------------------------
def build_cover(doc):
    # logo
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, 6, 6, 1.0)
    run = p.add_run()
    run.add_picture(str(FIG / "tmu_logo.jpeg"), width=Inches(1.15))

    add_text(doc, "TEERTHANKER MAHAVEER UNIVERSITY", size=16, bold=True, align="center",
             before=4, after=0, line=1.1, color=NAVY)
    add_text(doc, "MORADABAD (U.P.)", size=13, bold=True, align="center",
             before=0, after=4, line=1.1, color=NAVY)
    p = add_text(doc, "FACULTY OF ENGINEERING & COMPUTING SCIENCES", size=12, bold=True,
                 align="center", before=2, after=8, line=1.1, color=TEAL)
    add_hline(p, GOLD_HEX, "14")

    add_text(doc, "PROJECT REPORT", size=18, bold=True, align="center",
             before=16, after=2, line=1.1, color=NAVY)
    add_text(doc, "on", size=12, italic=True, align="center", before=2, after=6, line=1.1)
    add_text(doc, TITLE_LONG, size=16, bold=True, align="center",
             before=2, after=10, line=1.2, color=NAVY)

    p = add_text(doc, "", size=12, align="center", before=0, after=0, line=1.0)
    add_hline(p, NAVY_HEX, "8")

    add_text(doc, "Submitted in Partial Fulfilment of the Requirement for the Degree of",
             size=12, italic=True, align="center", before=10, after=4, line=1.15)
    add_text(doc, PROGRAM, size=14, bold=True, align="center", before=0, after=0, line=1.15, color=NAVY)
    add_text(doc, "(AI & Data Analytics)", size=12, italic=True, align="center", before=0, after=10, line=1.15)

    add_text(doc, "By", size=12, align="center", before=4, after=2, line=1.15)
    add_text(doc, f"{STUDENT.upper()}", size=16, bold=True, align="center",
             before=0, after=0, line=1.15, color=NAVY)
    add_text(doc, f"Enrollment No.: {ENROLL}", size=12, bold=True, align="center",
             before=0, after=2, line=1.15)
    add_text(doc, f"Semester – {SEMESTER}    |    Section – {SECTION}", size=12, align="center",
             before=0, after=2, line=1.15)
    add_text(doc, MONTH, size=12, align="center", before=4, after=12, line=1.15)

    # submitted to / by table
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left_head, right_head = table.rows[0].cells
    cell_text(left_head, "Submitted To", bold=True, size=12, align="center", fill=NAVY_HEX, color=WHITE)
    cell_text(right_head, "Submitted By", bold=True, size=12, align="center", fill=NAVY_HEX, color=WHITE)
    cell_text(table.rows[1].cells[0], f"Guide: {GUIDE1}\nGuide: {GUIDE2}", size=11, align="center", fill=SOFT_HEX)
    cell_text(table.rows[1].cells[1], f"{STUDENT}\n{ENROLL}", size=11, align="center", fill=SOFT_HEX)
    cell_text(table.rows[2].cells[0], FACULTY, size=11, align="center")
    cell_text(table.rows[2].cells[1], PROGRAM, size=11, align="center")
    cell_text(table.rows[3].cells[0], UNIV, size=11, align="center")
    cell_text(table.rows[3].cells[1], f"Section – {SECTION}    •    {MONTH}", size=11, align="center")
    style_table(table, header=False, col_widths=[3.2, 3.2])
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, top={"sz": "8", "color": NAVY_HEX},
                            bottom={"sz": "8", "color": NAVY_HEX},
                            left={"sz": "8", "color": NAVY_HEX},
                            right={"sz": "8", "color": NAVY_HEX})

    add_text(doc, "", size=10, after=4)
    p = add_text(doc, FACULTY.upper(), size=11, bold=True, align="center",
                 before=14, after=0, line=1.1, color=NAVY)
    add_text(doc, UNIV.upper(), size=11, bold=True, align="center",
             before=0, after=0, line=1.1, color=NAVY)


def build_declaration(doc):
    heading(doc, "DECLARATION", 1)
    p = add_text(doc, "", size=12, after=2)
    add_hline(p, GOLD_HEX, "12")

    add_text(
        doc,
        f"I, {STUDENT}, student of {PROGRAM}, Semester-{SEMESTER}, studying at "
        f"{UNIV} (U.P.), hereby declare that the Project Report / Dissertation entitled "
        f"“{TITLE_LONG}” submitted in partial fulfilment of the requirement for the degree of "
        f"BCA is an original work carried out by me.",
    )
    add_text(
        doc,
        "The information and analysis presented in this report are true to the best of my knowledge. "
        "This Project Report / Dissertation has not been submitted to any other University for the award "
        "of any other Degree, Diploma or Fellowship.",
    )
    add_text(
        doc,
        "I further declare that I have used open-source tools, publicly available datasets and a "
        "reference Streamlit implementation for academic demonstration, and that the documentation, "
        "system analysis, diagrams and report writing presented here are my own work for this submission.",
    )
    add_text(doc, f"Date: {DATE_STR}", align="left", before=16, after=2)
    add_text(doc, f"Place: {PLACE}", align="left", before=0, after=16)
    add_text(doc, "Submitted By:", bold=True, align="right", after=0)
    add_text(doc, f"{STUDENT}", bold=True, align="right", after=0, color=NAVY)
    add_text(doc, f"({ENROLL})", bold=True, align="right", after=0)


def build_certificate(doc):
    heading(doc, "CERTIFICATE", 1)
    p = add_text(doc, "", size=12, after=2)
    add_hline(p, GOLD_HEX, "12")

    add_text(
        doc,
        f"This is to certify that the Project Report / Dissertation entitled "
        f"“{TITLE_LONG}” has been submitted in partial fulfilment of the requirement for the degree of "
        f"BCA and has been carried out by the student whose name is listed below, under my / our "
        "supervision and guidance.",
    )
    add_text(doc, "Submitted By:", bold=True, before=12, after=2, align="left")
    add_text(doc, f"{STUDENT} ({ENROLL})", bold=True, size=13, after=12, color=NAVY)

    table = doc.add_table(rows=3, cols=2)
    cell_text(table.rows[0].cells[0], "Guide Name", bold=True, size=11, align="center", fill=NAVY_HEX, color=WHITE)
    cell_text(table.rows[0].cells[1], "Guide Signature", bold=True, size=11, align="center", fill=NAVY_HEX, color=WHITE)
    cell_text(table.rows[1].cells[0], GUIDE1, size=12, align="center", fill=SOFT_HEX)
    cell_text(table.rows[1].cells[1], "\n\n", size=12, align="center")
    cell_text(table.rows[2].cells[0], GUIDE2, size=12, align="center", fill=SOFT_HEX)
    cell_text(table.rows[2].cells[1], "\n\n", size=12, align="center")
    style_table(table, header=False, col_widths=[3.2, 3.2])

    add_text(doc, f"Date: {DATE_STR}", before=20, after=2)
    add_text(doc, f"Place: {PLACE}", after=2)
    add_text(
        doc,
        f"{FACULTY}\n{UNIV}",
        italic=True, align="left", before=12, after=4, color=MUTED,
    )


def build_ack(doc):
    heading(doc, "ACKNOWLEDGEMENT", 1)
    p = add_text(doc, "", size=12, after=2)
    add_hline(p, GOLD_HEX, "12")

    add_text(
        doc,
        f"I am thankful to all those who helped me complete this Project Report / Dissertation on "
        f"“{TITLE_LONG}”.",
    )
    add_text(
        doc,
        f"First, I express my sincere gratitude to my Supervisor / Guide {GUIDE2} and {GUIDE1}, "
        "who supported me from the beginning of this work. Their guidance, suggestions and academic "
        "direction helped me organise the problem statement, the implementation description and the "
        "final report.",
    )
    add_text(
        doc,
        "I also take this opportunity to express my deep sense of gratitude to the Hon’ble Principal "
        "and to the Faculty of Engineering & Computing Sciences, Teerthanker Mahaveer University, "
        "for providing an academic environment that made this work possible.",
    )
    add_text(
        doc,
        f"I express my sincere thanks to the Project / Dissertation Coordinator {GUIDE2}, TMU, for "
        "the interest taken and the help provided throughout the project. I am also grateful to my "
        "family and classmates for their encouragement during the preparation of this report.",
    )
    add_text(doc, f"Date: {DATE_STR}", bold=True, before=16, after=2)
    add_text(doc, f"{STUDENT} ({ENROLL})", bold=True, after=2, color=NAVY)


def build_abstract(doc):
    heading(doc, "ABSTRACT", 1)
    p = add_text(doc, "", size=12, after=2)
    add_hline(p, GOLD_HEX, "12")

    add_text(
        doc,
        "The Multiple Disease Prediction System is a machine-learning decision-support application "
        "developed to present three independent disease-prediction modules through one browser-based "
        "interface. Health-risk assessment often depends on disease-specific indicators, and a user "
        "would otherwise need separate tools, separate forms and separate model files for diabetes, "
        "heart disease and Parkinson’s disease. This project addresses that fragmentation by combining "
        "structured input forms, saved classification models and a consistent result display inside a "
        "single Streamlit application.",
    )
    add_text(
        doc,
        "Three classification pipelines are used. Diabetes prediction is performed with a Support "
        "Vector Machine (linear kernel) trained on the PIMA diabetes dataset (768 records, 8 clinical "
        "features). Heart-disease prediction is performed with Logistic Regression trained on a "
        "heart-disease dataset (303 records, 13 cardiac features). Parkinson’s prediction is performed "
        "with a Support Vector Machine (linear kernel) trained on a voice-measurement dataset "
        "(195 records, 22 acoustic features). The trained estimators are stored as pickle artifacts "
        "(diabetes_model.sav, heart_disease_model.sav and parkinsons_model.sav) and loaded once when "
        "the application starts.",
    )
    add_text(
        doc,
        "At runtime the user selects a module from the sidebar, enters the required numerical values, "
        "and submits the form. The application converts the inputs to float, builds an ordered feature "
        "vector and calls the corresponding model’s predict() method. A direct classification-style "
        "message is then shown. On the training notebooks accompanying this project, the diabetes model "
        "achieved about 77.27% test accuracy, the heart-disease model about 81.97% test accuracy, and "
        "the Parkinson’s model about 87.18% test accuracy.",
    )
    add_text(
        doc,
        "The system is implemented with Python, NumPy, scikit-learn, Streamlit and streamlit-option-menu. "
        "It is intended as an academic demonstration of end-to-end machine-learning deployment and as a "
        "prototype decision-support interface. It is not a clinical diagnostic device. The report "
        "documents the problem, scope, modules, methodology, architecture, technology stack, limitations "
        "and possible future enhancements.",
    )

    add_text(doc, "Keywords:", bold=True, before=10, after=2, align="left")
    add_text(
        doc,
        "Machine Learning; Streamlit; Disease Prediction; Support Vector Machine; Logistic Regression; "
        "Diabetes; Heart Disease; Parkinson’s Disease; Pickle; Decision Support.",
        italic=True, after=6,
    )


def build_toc(doc):
    heading(doc, "TABLE OF CONTENTS", 1)
    p = add_text(doc, "", size=12, after=2)
    add_hline(p, GOLD_HEX, "12")
    add_text(
        doc,
        "After opening this file in Microsoft Word, right-click the table below (or press Ctrl+A and F9) "
        "and choose Update Field if you need Word to refresh automatic page numbers. A complete manual "
        "contents list is provided for immediate reading and editing.",
        size=11, italic=True, color=MUTED, after=10, line=1.15,
    )

    entries = [
        ("DECLARATION", "i", True),
        ("CERTIFICATE", "ii", True),
        ("ACKNOWLEDGEMENT", "iii", True),
        ("ABSTRACT", "iv", True),
        ("TABLE OF CONTENTS", "v", True),
        ("LIST OF FIGURES", "vii", True),
        ("LIST OF TABLES", "viii", True),
        ("1.  PROJECT TITLE", "1", True),
        ("2.  PROBLEM STATEMENT", "1", True),
        ("3.  PROJECT DESCRIPTION", "2", True),
        ("     3.1  Aim, Objectives and Scope", "2", False),
        ("     3.2  Scope of the Work", "3", False),
        ("     3.3  Existing Approach versus Proposed System", "3", False),
        ("     3.4  Project Modules", "4", False),
        ("     3.5  Context Diagram (High Level)", "6", False),
        ("4.  IMPLEMENTATION METHODOLOGY", "7", True),
        ("     4.1  Sequential Machine-Learning Methodology", "7", False),
        ("     4.2  Data Collection and Datasets", "8", False),
        ("     4.3  Data Preprocessing", "9", False),
        ("     4.4  Feature Design by Disease Module", "10", False),
        ("     4.5  Model Training", "11", False),
        ("     4.6  Model Evaluation", "12", False),
        ("     4.7  Model Serialization", "13", False),
        ("     4.8  Application Inference Pipeline", "13", False),
        ("     4.9  Data Flow Diagram", "14", False),
        ("     4.10 Flow Chart", "15", False),
        ("     4.11 Use Case Diagram", "16", False),
        ("5.  SYSTEM ARCHITECTURE (HLD AND LLD)", "17", True),
        ("6.  DISEASE MODULES AND FEATURE DESIGN", "18", True),
        ("7.  USER INTERFACE AND NAVIGATION DESIGN", "20", True),
        ("8.  TECHNOLOGIES TO BE USED", "21", True),
        ("     8.1  Software Platform", "21", False),
        ("     8.2  Hardware Platform", "22", False),
        ("     8.3  Tools", "22", False),
        ("9.  FEASIBILITY ANALYSIS", "23", True),
        ("10. ADVANTAGES OF THIS PROJECT", "24", True),
        ("11. SECURITY, RELIABILITY AND LIMITATIONS", "25", True),
        ("12. ASSUMPTIONS", "26", True),
        ("13. EXPECTED IMPACT AND BENEFITS", "27", True),
        ("14. FUTURE SCOPE AND FURTHER ENHANCEMENT", "28", True),
        ("15. PROJECT REPOSITORY LOCATION", "29", True),
        ("16. DEFINITIONS, ACRONYMS AND ABBREVIATIONS", "30", True),
        ("17. CONCLUSION", "31", True),
        ("18. REFERENCES / BIBLIOGRAPHY", "32", True),
        ("APPENDIX A  Feature Dictionary", "34", True),
        ("APPENDIX B  Software Dependencies and Runtime Command", "36", True),
    ]
    table = doc.add_table(rows=len(entries), cols=2)
    table.autofit = False
    for i, (title, page, strong) in enumerate(entries):
        c0, c1 = table.rows[i].cells
        fill = LIGHT_HEX if i % 2 == 0 else "FFFFFF"
        cell_text(c0, title, bold=strong, size=11 if strong else 11, fill=fill,
                  color=NAVY if strong else BLACK)
        cell_text(c1, page, bold=strong, size=11, align="right", fill=fill,
                  color=NAVY if strong else BLACK)
        for edge in ("top", "left", "bottom", "right"):
            set_cell_border(c0, **{e: {"sz": "0", "color": "FFFFFF", "val": "nil"} for e in ("top", "left", "bottom", "right")})
            set_cell_border(c1, **{e: {"sz": "0", "color": "FFFFFF", "val": "nil"} for e in ("top", "left", "bottom", "right")})
        # light bottom line
        set_cell_border(c0, top={"sz": "0", "val": "nil", "color": "FFFFFF"},
                        left={"sz": "0", "val": "nil", "color": "FFFFFF"},
                        right={"sz": "0", "val": "nil", "color": "FFFFFF"},
                        bottom={"sz": "4", "color": "E2E8F0"})
        set_cell_border(c1, top={"sz": "0", "val": "nil", "color": "FFFFFF"},
                        left={"sz": "0", "val": "nil", "color": "FFFFFF"},
                        right={"sz": "0", "val": "nil", "color": "FFFFFF"},
                        bottom={"sz": "4", "color": "E2E8F0"})
        c0.width = Inches(5.5)
        c1.width = Inches(1.0)
    add_text(
        doc,
        "Note: Page numbers in this contents list are approximate. They will shift slightly after you "
        "edit the file or after Microsoft Word paginates the document on your computer. You may update "
        "the numbers freely.",
        size=10, italic=True, color=MUTED, before=10, line=1.15,
    )


def build_list_of_figures(doc):
    heading(doc, "LIST OF FIGURES", 1)
    p = add_text(doc, "", size=12, after=2)
    add_hline(p, GOLD_HEX, "12")
    figs = [
        ("Figure 1", "Context Diagram (Level 0) of the Multiple Disease Prediction System"),
        ("Figure 2", "Existing fragmented workflow versus the proposed unified system"),
        ("Figure 3", "Logical system architecture of the deployed application"),
        ("Figure 4", "End-to-end inference methodology used by the Streamlit application"),
        ("Figure 5", "Machine-learning training pipeline and runtime inference pipeline"),
        ("Figure 6", "Application flow chart"),
        ("Figure 7", "Use case diagram"),
        ("Figure 8", "High-level design (HLD) and low-level design (LLD)"),
        ("Figure 9", "User interface and sidebar navigation design"),
    ]
    table = doc.add_table(rows=len(figs) + 1, cols=2)
    cell_text(table.rows[0].cells[0], "Figure No.", bold=True, size=11, align="center", fill=NAVY_HEX, color=WHITE)
    cell_text(table.rows[0].cells[1], "Title", bold=True, size=11, align="center", fill=NAVY_HEX, color=WHITE)
    for i, (no, title) in enumerate(figs, start=1):
        fill = LIGHT_HEX if i % 2 else "FFFFFF"
        cell_text(table.rows[i].cells[0], no, size=11, align="center", fill=fill, color=NAVY)
        cell_text(table.rows[i].cells[1], title, size=11, fill=fill)
    style_table(table, header=False, col_widths=[1.4, 5.1])


def build_list_of_tables(doc):
    heading(doc, "LIST OF TABLES", 1)
    p = add_text(doc, "", size=12, after=2)
    add_hline(p, GOLD_HEX, "12")
    tabs = [
        ("Table 1", "Comparison of the three disease-prediction modules"),
        ("Table 2", "Diabetes module — input features"),
        ("Table 3", "Heart-disease module — input features"),
        ("Table 4", "Parkinson’s module — input features (voice measures)"),
        ("Table 5", "Datasets used for model training"),
        ("Table 6", "Algorithms, split settings and observed accuracy"),
        ("Table 7", "Software stack and role of each component"),
        ("Table 8", "Hardware requirements"),
        ("Table 9", "Feasibility summary"),
        ("Table 10", "Project artifacts and repository location"),
        ("Table 11", "Definitions, acronyms and abbreviations"),
        ("Table 12", "Application dependencies (requirements.txt)"),
    ]
    table = doc.add_table(rows=len(tabs) + 1, cols=2)
    cell_text(table.rows[0].cells[0], "Table No.", bold=True, size=11, align="center", fill=NAVY_HEX, color=WHITE)
    cell_text(table.rows[0].cells[1], "Title", bold=True, size=11, align="center", fill=NAVY_HEX, color=WHITE)
    for i, (no, title) in enumerate(tabs, start=1):
        fill = LIGHT_HEX if i % 2 else "FFFFFF"
        cell_text(table.rows[i].cells[0], no, size=11, align="center", fill=fill, color=NAVY)
        cell_text(table.rows[i].cells[1], title, size=11, fill=fill)
    style_table(table, header=False, col_widths=[1.4, 5.1])


# ---------------------------------------------------------------------------
# Body
# ---------------------------------------------------------------------------
def build_body(doc):
    heading(doc, "1.  PROJECT TITLE", 1)
    add_text(
        doc,
        "Design and Development of a Machine Learning–Based Multiple Disease Prediction System "
        "for Diabetes, Heart Disease and Parkinson’s Disease using a Unified Streamlit Interface",
        bold=True, size=13, align="center", color=NAVY, before=6, after=12,
    )
    add_text(
        doc,
        "The short title used throughout this report is Multiple Disease Prediction System. "
        "The implemented application presents three independent classification modules — Diabetes "
        "Prediction, Heart Disease Prediction and Parkinson’s Prediction — inside one Streamlit "
        "web interface titled “Health Assistant”.",
    )

    heading(doc, "2.  PROBLEM STATEMENT", 1)
    add_text(
        doc,
        "Health-risk assessment often involves different disease-specific indicators. Diabetes "
        "screening uses laboratory and demographic values such as glucose, BMI and age. Heart-disease "
        "assessment uses cardiac measurements such as chest-pain type, cholesterol, maximum heart rate "
        "and ST depression. Parkinson’s screening, in the dataset used here, depends on voice-related "
        "numerical measures such as jitter, shimmer, HNR and PPE. Because the feature sets are different, "
        "a user may need separate workflows, separate programs and separate model files for each task.",
    )
    add_text(
        doc,
        "Technical machine-learning models are also difficult to consume directly. A saved estimator "
        "in pickle format cannot be used by a non-technical user unless an interface collects inputs, "
        "converts them to numbers, calls predict(), and explains the output in ordinary language. "
        "Without that interface, even a correctly trained model remains inaccessible.",
    )
    add_text(
        doc,
        "The proposed project addresses these challenges by building a unified, browser-based "
        "prediction system. The application loads three saved classification models at startup, "
        "provides a sidebar for module selection, presents a structured input form for each disease, "
        "converts the entered values to a numeric feature vector, and displays a clear classification "
        "message. The same interaction pattern is used for all three diseases; only the input schema "
        "and the loaded model change.",
    )
    add_text(
        doc,
        "The project is positioned as an academic / prototype decision-support system. It demonstrates "
        "end-to-end machine-learning deployment — from dataset and training notebook to serialized "
        "artifact to user-facing inference — but it does not replace a clinical diagnosis, laboratory "
        "test or physician review.",
    )

    heading(doc, "3.  PROJECT DESCRIPTION", 1)
    add_text(
        doc,
        "The Multiple Disease Prediction System is a Python application that uses Streamlit for the "
        "user interface and scikit-learn estimators for inference. The source file app.py is the "
        "runtime program. It sets the page configuration, loads three model files from the "
        "saved_models directory, and renders one of three pages according to the sidebar choice.",
    )
    add_text(
        doc,
        "The three prediction tasks are independent. They do not share features, they do not share "
        "estimators, and a result on one page has no effect on another page. What they share is the "
        "application shell: wide layout, option-menu navigation, text-input forms, a test-result "
        "button, float conversion, model.predict([user_input]), and an st.success() message.",
    )

    heading(doc, "3.1  Aim, Objectives and Scope", 2)
    add_text(
        doc,
        "The aim of the project is to provide a single, easy-to-use interface through which a user "
        "can obtain machine-learning predictions for three diseases without interacting with model "
        "files, notebooks or source code.",
        before=4,
    )
    add_text(doc, "The specific objectives are:", bold=True, after=4, align="left")
    bullet(doc, "Provide disease-specific input forms for diabetes (8 features), heart disease (13 features) and Parkinson’s disease (22 features).")
    bullet(doc, "Convert user-entered values to numeric form before inference, using an ordered feature list that matches the training schema.")
    bullet(doc, "Connect each form to its corresponding saved model artifact loaded at application startup.")
    bullet(doc, "Display a direct classification-style diagnosis message immediately after the user submits the required values.")
    bullet(doc, "Keep navigation and result presentation consistent across modules so that the learning curve remains small.")

    add_text(doc, "The current scope is limited to the following:", bold=True, before=8, after=4, align="left")
    bullet(doc, "Diabetes prediction using 8 clinical / demographic inputs and diabetes_model.sav.")
    bullet(doc, "Heart-disease prediction using 13 cardiac inputs and heart_disease_model.sav.")
    bullet(doc, "Parkinson’s prediction using 22 voice-related numerical inputs and parkinsons_model.sav.")
    bullet(doc, "Inference performed with the loaded model’s predict() method; no retraining occurs inside the Streamlit app.")
    bullet(doc, "A prototype web UI. Authentication, patient databases, audit logs and hospital-system integration are outside the present scope.")

    heading(doc, "3.2  Scope of the Work", 2)
    add_text(
        doc,
        "The work covered by this project includes collection of public tabular datasets, exploratory "
        "inspection in training notebooks, train–test splitting, fitting of classification models, "
        "accuracy evaluation, serialization of the fitted estimators, and development of the Streamlit "
        "application that consumes those estimators. The major areas are:",
    )
    bullet(doc, "Use of three CSV datasets: diabetes.csv, heart.csv and parkinsons.csv.")
    bullet(doc, "Separation of features (X) and target labels (Y) for each disease.")
    bullet(doc, "An 80:20 train–test split, with stratification where the notebook applies it.")
    bullet(doc, "Training of SVM (linear kernel) for diabetes and Parkinson’s, and Logistic Regression for heart disease.")
    bullet(doc, "Evaluation using training accuracy and test accuracy.")
    bullet(doc, "Saving of models as .sav files with pickle.")
    bullet(doc, "Implementation of app.py with sidebar navigation, three input forms and result display.")
    bullet(doc, "Preparation of this project report, including architecture, diagrams, limitations and future work.")

    heading(doc, "3.3  Existing Approach versus Proposed System", 2)
    add_text(
        doc,
        "In a traditional or fragmented workflow, each disease has its own tool. The user must open "
        "a different interface, learn a different form, and interpret a different style of output. "
        "Repeated navigation and inconsistent interaction patterns increase the effort required even "
        "when the underlying task — enter numbers, obtain a class label — is the same.",
    )
    add_text(
        doc,
        "The proposed system places one Streamlit interface in front of three predictors. The user "
        "stays inside the same application, chooses a module from the sidebar, fills the corresponding "
        "form, and receives a prediction result in the same visual pattern. This reduces interaction "
        "complexity and makes the machine-learning models usable without exposing code.",
    )
    add_figure(
        doc, "fig2_existing_vs_proposed.png", 6.3,
        "Figure 2: Existing fragmented workflow versus the proposed unified Multiple Disease Prediction System.",
    )

    heading(doc, "3.4  Project Modules", 2)
    add_text(
        doc,
        "The project is organised into four major modules. Modules 1 to 3 correspond to the three "
        "disease pages in app.py. Module 4 covers training, evaluation and storage of the model artifacts "
        "that those pages consume.",
    )

    heading(doc, "Module 1: Diabetes Prediction", 3)
    add_text(
        doc,
        "This module collects eight inputs: Number of Pregnancies, Glucose Level, Blood Pressure, "
        "Skin Thickness, Insulin Level, BMI, Diabetes Pedigree Function and Age. The values are "
        "converted to float, passed to diabetes_model.sav, and mapped to one of two messages: "
        "“The person is diabetic” or “The person is not diabetic”. The underlying estimator is a "
        "Support Vector Classifier with a linear kernel, trained on the PIMA diabetes dataset with "
        "target column Outcome (0 = non-diabetic, 1 = diabetic).",
    )

    heading(doc, "Module 2: Heart Disease Prediction", 3)
    add_text(
        doc,
        "This module collects thirteen inputs used in standard heart-disease classification datasets: "
        "age, sex, chest-pain type, resting blood pressure, serum cholesterol, fasting blood sugar "
        "flag, resting electrocardiographic result, maximum heart rate, exercise-induced angina, "
        "ST depression (oldpeak), slope, number of major vessels (ca) and thal. The loaded model is "
        "a Logistic Regression estimator saved as heart_disease_model.sav. The output messages are "
        "“The person is having heart disease” or “The person does not have any heart disease”.",
    )

    heading(doc, "Module 3: Parkinson’s Disease Prediction", 3)
    add_text(
        doc,
        "This module collects twenty-two voice-related numerical features, including fundamental "
        "frequency measures (Fo, Fhi, Flo), jitter measures, RAP / PPQ / DDP, shimmer measures, "
        "NHR / HNR, RPDE, DFA, spread1, spread2, D2 and PPE. The name column from the original "
        "dataset is not used as a predictive feature. The estimator is an SVM with a linear kernel "
        "saved as parkinsons_model.sav. The output messages are “The person has Parkinson’s disease” "
        "or “The person does not have Parkinson’s disease”.",
    )

    heading(doc, "Module 4: Model Training, Evaluation and Storage", 3)
    add_text(
        doc,
        "Training is performed in Jupyter / Colab notebooks, not inside the Streamlit process. Each "
        "notebook loads a CSV file, inspects shape and class balance, splits features and target, "
        "fits the chosen estimator, prints training and test accuracy, demonstrates a sample "
        "prediction, and writes the fitted object to a .sav file with pickle.dump(). The application "
        "later reads those files with pickle.load() at startup. This separation keeps the user "
        "interface light and avoids retraining on every request.",
    )

    table = doc.add_table(rows=5, cols=4)
    heads = ["Aspect", "Diabetes", "Heart Disease", "Parkinson’s"]
    for i, h in enumerate(heads):
        cell_text(table.rows[0].cells[i], h, bold=True, size=10, align="center", fill=NAVY_HEX, color=WHITE)
    rows = [
        ["Input count", "8", "13", "22"],
        ["Input type", "Clinical / demographic", "Clinical / cardiac", "Voice-related numerical"],
        ["Model call", "predict([x])", "predict([x])", "predict([x])"],
        ["Result", "Diabetic / not diabetic", "Heart disease / no heart disease", "Parkinson’s / no Parkinson’s"],
    ]
    for r, row in enumerate(rows, start=1):
        fill = LIGHT_HEX if r % 2 else "FFFFFF"
        for c, val in enumerate(row):
            cell_text(table.rows[r].cells[c], val, size=10, fill=fill,
                      bold=(c == 0), align="center" if c else "left")
    style_table(table, header=False, col_widths=[1.5, 1.7, 1.7, 1.7])
    caption(doc, "Table 1: Comparison of the three disease-prediction modules.")

    heading(doc, "3.5  Context Diagram (High Level)", 2)
    add_text(
        doc,
        "At the context level the system has one external actor — the user — and one process — the "
        "Multiple Disease Prediction App. The user supplies disease-specific input values. The app "
        "converts those values into a feature vector and sends it to the disease-specific model. The "
        "model returns a class label, which the app presents as a readable result. The same context "
        "flow applies to all three diseases; only the feature set and the loaded artifact change.",
    )
    add_figure(
        doc, "fig1_context.png", 6.3,
        "Figure 1: Context-level (Level 0) data flow of the Multiple Disease Prediction System.",
    )

    heading(doc, "4.  IMPLEMENTATION METHODOLOGY", 1)
    add_text(
        doc,
        "Implementation follows a sequential machine-learning methodology for training, and a "
        "fixed inference methodology for the deployed application. Training is performed once. "
        "Inference is performed every time the user clicks a test-result button.",
    )

    heading(doc, "4.1  Sequential Machine-Learning Methodology", 2)
    add_text(doc, "The training and deployment sequence is as follows.", after=4)
    numbered(doc, "Data collection — obtain diabetes.csv, heart.csv and parkinsons.csv.")
    numbered(doc, "Data inspection — review shape, column names, missing values and class distribution.")
    numbered(doc, "Feature / label separation — drop the target column (and the name column for Parkinson’s) to form X; keep the label as Y.")
    numbered(doc, "Train–test split — 80% training and 20% testing, with a fixed random_state for reproducibility; stratification is used for diabetes and heart disease.")
    numbered(doc, "Model training — fit SVC(kernel='linear') for diabetes and Parkinson’s; fit LogisticRegression() for heart disease.")
    numbered(doc, "Model evaluation — compute accuracy_score on both the training set and the test set.")
    numbered(doc, "Sample prediction — run a hand-crafted input through the fitted model to confirm the output mapping.")
    numbered(doc, "Model serialization — pickle.dump() the fitted estimator to a .sav file.")
    numbered(doc, "Application assembly — write app.py, load the three artifacts, and wire each form to predict().")
    numbered(doc, "Interactive inference — convert form values to float, call predict(), and display a diagnosis message.")

    add_figure(
        doc, "fig4_methodology.png", 6.3,
        "Figure 4: End-to-end inference methodology implemented in the Streamlit application.",
    )

    heading(doc, "4.2  Data Collection and Datasets", 2)
    add_text(
        doc,
        "The project uses three publicly known tabular datasets that match the feature lists in app.py. "
        "The README of the source project states that the training notebooks and the datasets are "
        "provided in their respective folders. The dataset files are diabetes.csv, heart.csv and "
        "parkinsons.csv.",
    )

    table = doc.add_table(rows=4, cols=5)
    heads = ["Dataset file", "Common source", "Records", "Predictive features", "Target"]
    for i, h in enumerate(heads):
        cell_text(table.rows[0].cells[i], h, bold=True, size=9, align="center", fill=NAVY_HEX, color=WHITE)
    data = [
        ["diabetes.csv", "PIMA Indians Diabetes", "768", "8", "Outcome (0/1)"],
        ["heart.csv", "UCI / Kaggle Heart Disease", "303", "13", "target (0/1)"],
        ["parkinsons.csv", "UCI Parkinson’s (voice)", "195", "22 (name dropped)", "status (0/1)"],
    ]
    for r, row in enumerate(data, start=1):
        fill = LIGHT_HEX if r % 2 else "FFFFFF"
        for c, val in enumerate(row):
            cell_text(table.rows[r].cells[c], val, size=9, fill=fill, align="center")
    style_table(table, header=False, col_widths=[1.3, 1.7, 0.9, 1.5, 1.2])
    caption(doc, "Table 5: Datasets used for model training (record counts as observed in the training notebooks).")

    add_text(
        doc,
        "The PIMA diabetes dataset contains clinical measurements commonly used in educational "
        "machine-learning courses. The heart-disease dataset contains 14 columns including the target. "
        "The Parkinson’s dataset contains 24 columns in total; the patient name is an identifier and "
        "is removed before training, leaving 22 acoustic features plus the binary status label. In "
        "the Parkinson’s notebook the class distribution is 147 positive (status = 1) and 48 negative "
        "(status = 0), so the dataset is imbalanced. This is a known limitation of the training set.",
    )

    heading(doc, "4.3  Data Preprocessing", 2)
    add_text(
        doc,
        "Preprocessing in the training notebooks is deliberately simple and matches a first-course "
        "machine-learning workflow:",
    )
    bullet(doc, "Load each CSV into a pandas DataFrame.")
    bullet(doc, "Inspect the first rows, data types, descriptive statistics and target value counts.")
    bullet(doc, "Confirm missing-value counts (the heart and Parkinson’s notebooks report no nulls).")
    bullet(doc, "Separate X and Y. For Parkinson’s, drop both name and status from X.")
    bullet(doc, "Apply sklearn.model_selection.train_test_split with test_size=0.2 and random_state=2.")
    bullet(doc, "For diabetes and heart disease, set stratify=Y so that class proportions are preserved in the split.")
    add_text(
        doc,
        "No StandardScaler, imputation pipeline or explicit outlier treatment is applied in the "
        "notebooks that produce the saved models. The Streamlit app likewise performs only float() "
        "conversion. This keeps the student implementation short, but it also means that blank fields, "
        "non-numeric text or values far outside the training range can cause runtime errors or "
        "unreliable predictions. These points are treated as limitations and as future-work items.",
    )

    heading(doc, "4.4  Feature Design by Disease Module", 2)
    add_text(
        doc,
        "Feature order is part of the contract between the training notebooks and app.py. The list "
        "passed to predict() must follow the column order of X. The application therefore builds "
        "user_input as a Python list in a fixed sequence, then casts every element to float.",
    )
    add_text(
        doc,
        "Diabetes uses eight clinical features. Heart disease uses thirteen mixed numeric and encoded "
        "categorical features (for example sex, chest-pain type, fasting blood sugar, slope, ca and "
        "thal are stored as numbers). Parkinson’s uses twenty-two acoustic measures derived from "
        "voice recordings. The complete dictionaries are given in Chapter 6 and in Appendix A.",
    )

    heading(doc, "4.5  Model Training", 2)
    add_text(
        doc,
        "Each disease has its own notebook and its own estimator. The choices below are those used "
        "in the project’s training notebooks, not invented for this report.",
    )

    heading(doc, "Diabetes — Support Vector Machine", 3)
    add_text(
        doc,
        "The diabetes notebook imports svm from sklearn and creates classifier = svm.SVC(kernel='linear'). "
        "The linear kernel is appropriate for a moderately sized tabular problem and yields a "
        "maximum-margin linear decision boundary in the eight-dimensional feature space. After the "
        "80:20 stratified split the training matrix has shape (614, 8) and the test matrix has shape "
        "(154, 8). The fitted classifier is written to diabetes_model.sav.",
    )

    heading(doc, "Heart disease — Logistic Regression", 3)
    add_text(
        doc,
        "The heart notebook imports LogisticRegression from sklearn.linear_model and calls model.fit() "
        "on the training split. Logistic Regression estimates the probability of the positive class "
        "and is a standard baseline for this dataset. After the stratified 80:20 split the training "
        "matrix has shape (242, 13) and the test matrix has shape (61, 13). The original notebook "
        "records a convergence warning from the lbfgs solver (iteration limit reached). The fitted "
        "model is nevertheless saved as heart_disease_model.sav and is the artifact loaded by app.py. "
        "Increasing max_iter is a natural future improvement.",
    )

    heading(doc, "Parkinson’s — Support Vector Machine", 3)
    add_text(
        doc,
        "The Parkinson’s notebook also uses svm.SVC(kernel='linear'). After dropping name and status, "
        "X has 22 columns. The 80:20 split (random_state=2, without stratify in the notebook) produces "
        "X_train of shape (156, 22) and X_test of shape (39, 22). The fitted model is saved as "
        "parkinsons_model.sav.",
    )
    add_figure(
        doc, "fig5_ml_pipeline.png", 6.3,
        "Figure 5: Training pipeline (notebooks) and inference pipeline (app.py).",
    )

    heading(doc, "4.6  Model Evaluation", 2)
    add_text(
        doc,
        "The notebooks evaluate models with sklearn.metrics.accuracy_score on the training set and "
        "on the held-out test set. Precision, recall, F1-score, ROC-AUC and confusion matrices are "
        "not printed in those notebooks, and app.py does not display metrics to the end user. The "
        "accuracy figures below are taken from the notebook outputs so that this report does not "
        "invent validation statistics.",
    )

    table = doc.add_table(rows=4, cols=6)
    heads = ["Disease", "Algorithm", "Train rows", "Test rows", "Train accuracy", "Test accuracy"]
    for i, h in enumerate(heads):
        cell_text(table.rows[0].cells[i], h, bold=True, size=9, align="center", fill=NAVY_HEX, color=WHITE)
    data = [
        ["Diabetes", "SVM (linear)", "614", "154", "78.34%", "77.27%"],
        ["Heart disease", "Logistic Regression", "242", "61", "85.12%", "81.97%"],
        ["Parkinson’s", "SVM (linear)", "156", "39", "87.18%", "87.18%"],
    ]
    for r, row in enumerate(data, start=1):
        fill = LIGHT_HEX if r % 2 else "FFFFFF"
        for c, val in enumerate(row):
            cell_text(table.rows[r].cells[c], val, size=9, fill=fill, align="center")
    style_table(table, header=False, col_widths=[1.15, 1.35, 0.9, 0.85, 1.2, 1.15])
    caption(doc, "Table 6: Algorithms, split sizes and accuracy scores reported by the training notebooks.")

    add_text(
        doc,
        "Interpretation should remain cautious. Accuracy on a small and, in the Parkinson’s case, "
        "imbalanced test set can overstate operational quality. Equal train and test accuracy for "
        "Parkinson’s (both 87.18%) is possible on 39 test rows but is not by itself proof of "
        "robust generalisation. The Streamlit result is a class label, not a calibrated probability "
        "and not a medical diagnosis.",
    )

    heading(doc, "4.7  Model Serialization", 2)
    add_text(
        doc,
        "Each notebook saves its estimator with pickle:",
    )
    add_text(doc, "filename = 'diabetes_model.sav'          # similarly heart_disease_model.sav, parkinsons_model.sav",
             font="Courier New", size=10, align="left", line=1.15, before=4, after=4)
    add_text(doc, "pickle.dump(model, open(filename, 'wb'))",
             font="Courier New", size=10, align="left", line=1.15, before=0, after=8)
    add_text(
        doc,
        "The application reloads the artifacts at import / startup time using paths relative to the "
        "directory of app.py:",
    )
    add_text(
        doc,
        "working_dir = os.path.dirname(os.path.abspath(__file__))\n"
        "diabetes_model = pickle.load(open(f'{working_dir}/saved_models/diabetes_model.sav', 'rb'))\n"
        "heart_disease_model = pickle.load(open(f'{working_dir}/saved_models/heart_disease_model.sav', 'rb'))\n"
        "parkinsons_model = pickle.load(open(f'{working_dir}/saved_models/parkinsons_model.sav', 'rb'))",
        font="Courier New", size=10, align="left", line=1.15, before=4, after=8,
    )
    add_text(
        doc,
        "Loading at startup means that a missing or corrupt .sav file will fail before any page is "
        "used. It also means that the three models remain in memory for the life of the Streamlit "
        "session, which is appropriate for small scikit-learn classifiers.",
    )

    heading(doc, "4.8  Application Inference Pipeline", 2)
    add_text(
        doc,
        "Every disease page follows the same low-level sequence. This is the actual control flow "
        "implemented in app.py.",
    )
    numbered(doc, "Render the page title and a multi-column layout of st.text_input widgets.")
    numbered(doc, "Initialise an empty diagnosis string.")
    numbered(doc, "If the user clicks the test-result button, collect widget values into a Python list in training order.")
    numbered(doc, "Convert the list with user_input = [float(x) for x in user_input].")
    numbered(doc, "Call the loaded estimator: prediction = model.predict([user_input]).")
    numbered(doc, "If prediction[0] == 1, set a positive diagnosis message; otherwise set a negative message.")
    numbered(doc, "Display the message with st.success(diagnosis).")
    add_text(
        doc,
        "Because st.success() sits outside the button block, an empty green box is visible before "
        "the first click. After a successful click the message appears. Invalid or blank input raises "
        "a Python exception during float conversion; the present code does not catch that exception. "
        "Range checks and try/except handling are listed under future enhancements.",
    )

    heading(doc, "4.9  Data Flow Diagram", 2)
    add_text(
        doc,
        "Data movement at runtime is one-way from the user to the model and back as a label. No "
        "database write occurs. No training data is read by app.py. The only persistent files "
        "involved at inference time are the three .sav artifacts.",
    )
    add_figure(
        doc, "fig3_architecture.png", 6.3,
        "Figure 3: Logical architecture — user, Streamlit UI, input cast, model artifacts, prediction and result display.",
    )

    heading(doc, "4.10  Flow Chart", 2)
    add_text(
        doc,
        "The flow chart records the operational path of the running application, from process start "
        "through model loading, sidebar selection, data entry, the test-result decision, prediction "
        "and message display. If the button is not clicked, the application simply continues to show "
        "the form.",
    )
    add_figure(
        doc, "fig6_flowchart.png", 4.6,
        "Figure 6: Application flow chart of the Multiple Disease Prediction System.",
    )

    heading(doc, "4.11  Use Case Diagram", 2)
    add_text(
        doc,
        "The primary actor is a health user or student tester. The actor can select a disease module, "
        "enter the required inputs, request a prediction for diabetes, heart disease or Parkinson’s, "
        "and view the diagnosis message. Loading of saved models and execution of predict() are "
        "system-side use cases triggered by those requests.",
    )
    add_figure(
        doc, "fig7_usecase.png", 6.2,
        "Figure 7: Use case diagram for the Multiple Disease Prediction System.",
    )
    add_text(doc, "Main use cases:", bold=True, after=4, align="left")
    bullet(doc, "Select a disease prediction module from the sidebar.")
    bullet(doc, "Enter clinical or voice-related numerical information.")
    bullet(doc, "Request a diabetes, heart-disease or Parkinson’s prediction.")
    bullet(doc, "Review the classification result. The result is a screening aid, not a prescription for treatment.")

    heading(doc, "5.  SYSTEM ARCHITECTURE (HLD AND LLD)", 1)
    add_text(
        doc,
        "The high-level design (HLD) describes four logical layers. The low-level design (LLD) maps "
        "those layers onto the concrete calls found in app.py.",
    )
    add_text(doc, "High-level design layers:", bold=True, after=4, align="left")
    bullet(doc, "UI / Navigation — Streamlit page configuration, wide layout, and streamlit-option-menu sidebar.", bold_prefix="")
    bullet(doc, "Disease modules — three independent pages, each with its own widgets and feature list.")
    bullet(doc, "Model layer — pickle-loaded estimators stored as .sav files and held in memory.")
    bullet(doc, "Result layer — mapping of class label 0 / 1 to an English sentence shown with st.success().")
    add_text(doc, "Low-level design mapping:", bold=True, before=8, after=4, align="left")
    bullet(doc, "Input widgets — st.text_input() for every feature.")
    bullet(doc, "Feature list — ordered user_input list that matches training column order.")
    bullet(doc, "Conversion — float(x) applied to each list element.")
    bullet(doc, "Inference — model.predict([user_input]).")
    bullet(doc, "Output — st.success(message).")
    add_figure(
        doc, "fig8_hld_lld.png", 6.3,
        "Figure 8: High-level design layers and the repeating low-level implementation pattern.",
    )

    heading(doc, "6.  DISEASE MODULES AND FEATURE DESIGN", 1)
    add_text(
        doc,
        "This chapter lists the features exactly as they appear in the Streamlit forms and in the "
        "training matrices. Keeping this list accurate is essential, because a permutation of inputs "
        "would silently produce a wrong prediction.",
    )

    heading(doc, "6.1  Diabetes — eight features", 2)
    table = doc.add_table(rows=9, cols=3)
    for i, h in enumerate(["S.No.", "Form label in app.py", "Meaning / role"]):
        cell_text(table.rows[0].cells[i], h, bold=True, size=10, align="center", fill=NAVY_HEX, color=WHITE)
    diab = [
        ("1", "Number of Pregnancies", "Count of pregnancies (PIMA feature Pregnancies)"),
        ("2", "Glucose Level", "Plasma glucose concentration"),
        ("3", "Blood Pressure value", "Diastolic blood pressure"),
        ("4", "Skin Thickness value", "Triceps skin-fold thickness"),
        ("5", "Insulin Level", "2-hour serum insulin"),
        ("6", "BMI value", "Body mass index"),
        ("7", "Diabetes Pedigree Function value", "Family-history based pedigree function"),
        ("8", "Age of the Person", "Age in years"),
    ]
    for r, row in enumerate(diab, start=1):
        fill = LIGHT_HEX if r % 2 else "FFFFFF"
        for c, val in enumerate(row):
            cell_text(table.rows[r].cells[c], val, size=10, fill=fill, align="center" if c == 0 else "left")
    style_table(table, header=False, col_widths=[0.7, 2.6, 3.2])
    caption(doc, "Table 2: Diabetes module — input features in the order sent to diabetes_model.predict().")

    heading(doc, "6.2  Heart disease — thirteen features", 2)
    table = doc.add_table(rows=14, cols=3)
    for i, h in enumerate(["S.No.", "Form label in app.py", "Meaning / role"]):
        cell_text(table.rows[0].cells[i], h, bold=True, size=10, align="center", fill=NAVY_HEX, color=WHITE)
    heart = [
        ("1", "Age", "Age in years"),
        ("2", "Sex", "Encoded sex (dataset coding, typically 1 = male, 0 = female)"),
        ("3", "Chest Pain types", "Chest-pain type code (cp)"),
        ("4", "Resting Blood Pressure", "Resting blood pressure (trestbps)"),
        ("5", "Serum Cholestoral in mg/dl", "Serum cholesterol (chol)"),
        ("6", "Fasting Blood Sugar > 120 mg/dl", "Binary fasting blood-sugar flag (fbs)"),
        ("7", "Resting Electrocardiographic results", "Resting ECG result code (restecg)"),
        ("8", "Maximum Heart Rate achieved", "thalach"),
        ("9", "Exercise Induced Angina", "Binary flag (exang)"),
        ("10", "ST depression induced by exercise", "oldpeak"),
        ("11", "Slope of the peak exercise ST segment", "slope"),
        ("12", "Major vessels colored by flourosopy", "ca (0–3 / 4 depending on coding)"),
        ("13", "thal: 0 = normal; 1 = fixed defect; 2 = reversable defect", "thal code as labelled in the form"),
    ]
    for r, row in enumerate(heart, start=1):
        fill = LIGHT_HEX if r % 2 else "FFFFFF"
        for c, val in enumerate(row):
            cell_text(table.rows[r].cells[c], val, size=10, fill=fill, align="center" if c == 0 else "left")
    style_table(table, header=False, col_widths=[0.7, 2.8, 3.0])
    caption(doc, "Table 3: Heart-disease module — input features in the order sent to heart_disease_model.predict().")

    heading(doc, "6.3  Parkinson’s — twenty-two voice features", 2)
    add_text(
        doc,
        "Parkinson’s page widgets are arranged in five columns. The names follow the MDVP / acoustic "
        "columns of the UCI Parkinson’s dataset.",
    )
    table = doc.add_table(rows=23, cols=3)
    for i, h in enumerate(["S.No.", "Form label in app.py", "Group"]):
        cell_text(table.rows[0].cells[i], h, bold=True, size=10, align="center", fill=NAVY_HEX, color=WHITE)
    park = [
        ("1", "MDVP:Fo(Hz)", "Fundamental frequency"),
        ("2", "MDVP:Fhi(Hz)", "Fundamental frequency"),
        ("3", "MDVP:Flo(Hz)", "Fundamental frequency"),
        ("4", "MDVP:Jitter(%)", "Jitter"),
        ("5", "MDVP:Jitter(Abs)", "Jitter"),
        ("6", "MDVP:RAP", "Jitter"),
        ("7", "MDVP:PPQ", "Jitter"),
        ("8", "Jitter:DDP", "Jitter"),
        ("9", "MDVP:Shimmer", "Shimmer"),
        ("10", "MDVP:Shimmer(dB)", "Shimmer"),
        ("11", "Shimmer:APQ3", "Shimmer"),
        ("12", "Shimmer:APQ5", "Shimmer"),
        ("13", "MDVP:APQ", "Shimmer"),
        ("14", "Shimmer:DDA", "Shimmer"),
        ("15", "NHR", "Noise-to-harmonics"),
        ("16", "HNR", "Harmonics-to-noise"),
        ("17", "RPDE", "Nonlinear dynamical complexity"),
        ("18", "DFA", "Signal-fractal scaling"),
        ("19", "spread1", "Nonlinear frequency variation"),
        ("20", "spread2", "Nonlinear frequency variation"),
        ("21", "D2", "Dynamical complexity"),
        ("22", "PPE", "Pitch-period entropy"),
    ]
    for r, row in enumerate(park, start=1):
        fill = LIGHT_HEX if r % 2 else "FFFFFF"
        for c, val in enumerate(row):
            cell_text(table.rows[r].cells[c], val, size=10, fill=fill, align="center" if c == 0 else "left")
    style_table(table, header=False, col_widths=[0.7, 2.4, 3.4])
    caption(doc, "Table 4: Parkinson’s module — 22 voice-related features in predict() order.")

    heading(doc, "7.  USER INTERFACE AND NAVIGATION DESIGN", 1)
    add_text(
        doc,
        "The interface is a Streamlit web page with layout='wide' and page title “Health Assistant”. "
        "Navigation is implemented with streamlit_option_menu.option_menu inside st.sidebar. The menu "
        "title is “Multiple Disease Prediction System”. The three items are Diabetes Prediction, "
        "Heart Disease Prediction and Parkinsons Prediction, with icons activity, heart and person "
        "and a hospital-fill menu icon. default_index is 0, so the diabetes page opens first.",
    )
    add_text(
        doc,
        "Each disease page uses a multi-column grid of text inputs so that a large feature set still "
        "fits on one screen. Diabetes and heart disease use three columns; Parkinson’s uses five. "
        "A single primary button on each page triggers inference. Design strengths of the present UI "
        "are: wide page layout, disease-specific forms, a consistent submit action, immediate textual "
        "output, and a small learning curve.",
    )
    add_figure(
        doc, "fig9_ui.png", 6.3,
        "Figure 9: Sidebar navigation and the diabetes input form with test-result action and message area.",
    )
    add_text(
        doc,
        "UI limitations that should be improved later include the use of free-text fields instead of "
        "number inputs or select boxes for coded features (sex, chest-pain type, thal, and so on), "
        "the absence of placeholder examples, the absence of allowed ranges, and the empty success "
        "box shown before the first prediction.",
    )

    heading(doc, "8.  TECHNOLOGIES TO BE USED", 1)
    add_text(
        doc,
        "The project uses a lightweight Python stack. No GPU, no database server and no separate "
        "front-end framework are required for the implemented application.",
    )

    heading(doc, "8.1  Software Platform", 2)
    add_text(doc, "a) Front-end", bold=True, after=4, align="left")
    bullet(doc, " Used to develop the interactive web interface, page layout, input forms, buttons and result display.",
           bold_prefix="Streamlit: ")
    bullet(doc, " Used to render the sidebar navigation among the three disease modules.",
           bold_prefix="streamlit-option-menu: ")
    add_text(doc, "b) Back-end / machine learning", bold=True, before=8, after=4, align="left")
    bullet(doc, " Core programming language of app.py and of the training notebooks.",
           bold_prefix="Python: ")
    bullet(doc, " Numerical arrays used when a feature vector is passed to predict().",
           bold_prefix="NumPy: ")
    bullet(doc, " SVM and Logistic Regression estimators, train_test_split and accuracy_score.",
           bold_prefix="scikit-learn: ")
    bullet(doc, " Used in the notebooks for CSV loading, inspection and train/test frames.",
           bold_prefix="Pandas: ")
    bullet(doc, " Serialization and loading of .sav model artifacts.",
           bold_prefix="pickle: ")
    bullet(doc, " Used for model development, evaluation and experiment tracking.",
           bold_prefix="Jupyter Notebook / Google Colab: ")

    table = doc.add_table(rows=7, cols=3)
    for i, h in enumerate(["Component", "Technology", "Role in this project"]):
        cell_text(table.rows[0].cells[i], h, bold=True, size=10, align="center", fill=NAVY_HEX, color=WHITE)
    stack = [
        ["Application language", "Python 3", "Implements app.py and training notebooks"],
        ["Web UI / runtime", "Streamlit 1.29.0", "Browser forms, layout, result widgets"],
        ["Sidebar navigation", "streamlit-option-menu 0.3.6", "Module switching"],
        ["ML library", "scikit-learn 1.3.2", "SVM, Logistic Regression, metrics, split"],
        ["Numerical computing", "NumPy 1.26.3", "Numeric feature vectors"],
        ["Model files", "pickle .sav artifacts", "Persist and reload trained estimators"],
    ]
    for r, row in enumerate(stack, start=1):
        fill = LIGHT_HEX if r % 2 else "FFFFFF"
        for c, val in enumerate(row):
            cell_text(table.rows[r].cells[c], val, size=10, fill=fill)
    style_table(table, header=False, col_widths=[1.7, 2.0, 2.8])
    caption(doc, "Table 7: Software stack and the role of each component.")

    heading(doc, "8.2  Hardware Platform", 2)
    add_text(
        doc,
        "The project can be developed and executed on a standard personal computer or laptop capable "
        "of running a Python environment. Specialised GPU hardware is not required, because the "
        "models are conventional scikit-learn classifiers on small tabular datasets.",
    )
    table = doc.add_table(rows=6, cols=2)
    cell_text(table.rows[0].cells[0], "Item", bold=True, size=11, align="center", fill=NAVY_HEX, color=WHITE)
    cell_text(table.rows[0].cells[1], "Recommended specification", bold=True, size=11, align="center", fill=NAVY_HEX, color=WHITE)
    hw = [
        ("Processor", "Intel Core i3 / i5 or equivalent"),
        ("RAM", "Minimum 8 GB"),
        ("Storage", "Minimum 10 GB of available space"),
        ("Operating system", "Windows, Linux or macOS"),
        ("Internet", "Required when installing Python packages"),
    ]
    for r, (a, b) in enumerate(hw, start=1):
        fill = LIGHT_HEX if r % 2 else "FFFFFF"
        cell_text(table.rows[r].cells[0], a, size=11, fill=fill, bold=True)
        cell_text(table.rows[r].cells[1], b, size=11, fill=fill)
    style_table(table, header=False, col_widths=[2.2, 4.3])
    caption(doc, "Table 8: Hardware platform recommended for development and demonstration.")

    heading(doc, "8.3  Tools", 2)
    add_text(doc, "The following tools may be used during implementation and documentation:", after=4)
    bullet(doc, "for version control.", bold_prefix="Git ")
    bullet(doc, "for source-code repository management.", bold_prefix="GitHub ")
    bullet(doc, "for editing app.py and the report.", bold_prefix="Visual Studio Code / any Python IDE ")
    bullet(doc, "for training and evaluation.", bold_prefix="Jupyter Notebook / Google Colab ")
    bullet(doc, "for isolated installation of requirements.txt.", bold_prefix="Python virtual environment ")
    bullet(doc, "Microsoft Word, for editing this project report after generation.", bold_prefix="Office suite ")

    heading(doc, "9.  FEASIBILITY ANALYSIS", 1)
    add_text(
        doc,
        "Feasibility is examined under four headings that are commonly required in a BCA project "
        "synopsis: technical, operational, economic and scalability.",
    )
    table = doc.add_table(rows=5, cols=3)
    for i, h in enumerate(["Dimension", "Rating", "Justification based on the implemented system"]):
        cell_text(table.rows[0].cells[i], h, bold=True, size=10, align="center", fill=NAVY_HEX, color=WHITE)
    feas = [
        ["Technical", "High", "Python and Streamlit provide a lightweight web layer; saved model artifacts are loaded directly; no specialised hardware is required."],
        ["Operational", "High", "Users interact through forms and a sidebar. They never handle .sav files, notebooks or predict() calls themselves."],
        ["Economic", "High", "The application is built from common open-source Python packages and local model files. There is no licence cost for the implemented stack."],
        ["Scalability", "Moderate", "The modular structure can be extended with more diseases, but production scaling, authentication, monitoring and model serving would need extra engineering."],
    ]
    for r, row in enumerate(feas, start=1):
        fill = LIGHT_HEX if r % 2 else "FFFFFF"
        cell_text(table.rows[r].cells[0], row[0], size=10, fill=fill, bold=True, align="center")
        cell_text(table.rows[r].cells[1], row[1], size=10, fill=fill, bold=True, align="center", color=TEAL)
        cell_text(table.rows[r].cells[2], row[2], size=10, fill=fill)
    style_table(table, header=False, col_widths=[1.3, 1.1, 4.1])
    caption(doc, "Table 9: Feasibility summary of the Multiple Disease Prediction System.")

    heading(doc, "10.  ADVANTAGES OF THIS PROJECT", 1)
    add_text(
        doc,
        "The proposed system provides several practical advantages for an academic machine-learning "
        "deployment and for a first-level screening demonstration.",
    )
    bullet(doc, "One application replaces three disconnected tools. The user learns a single navigation pattern.",
           bold_prefix="Unified access. ")
    bullet(doc, "Sidebar choice plus a standard form-and-button pattern keeps interaction consistent.",
           bold_prefix="Standardised interaction. ")
    bullet(doc, "A non-programmer can obtain a model output without opening notebooks or writing Python.",
           bold_prefix="Accessible inference. ")
    bullet(doc, "Saved .sav files are reused on every run; the app does not retrain.",
           bold_prefix="Reusable models. ")
    bullet(doc, "The project shows the full path from CSV dataset to trained estimator to browser UI, which is a core BCA / ML learning outcome.",
           bold_prefix="End-to-end demonstration. ")
    bullet(doc, "A fourth disease can be added by training a new artifact and inserting one more sidebar page that follows the same LLD pattern.",
           bold_prefix="Modular extension. ")
    bullet(doc, "Python, Streamlit and scikit-learn are free and well documented.",
           bold_prefix="Low cost. ")

    heading(doc, "11.  SECURITY, RELIABILITY AND LIMITATIONS", 1)
    add_text(doc, "Current good practice reflected in the project:", bold=True, after=4, align="left")
    bullet(doc, "Keep model files in a controlled application directory (saved_models/) rather than allowing users to upload arbitrary pickle files.")
    bullet(doc, "Do not expose model internals, coefficients or serialized bytes in the user interface.")
    bullet(doc, "Convert inputs to numeric form before they reach predict(), so that the estimator receives the expected type.")
    add_text(doc, "Limitations identified from the implemented source:", bold=True, before=10, after=4, align="left")
    bullet(doc, "app.py does not contain training code. Training details in this report are taken from the accompanying notebooks and datasets, not from the UI file alone.")
    bullet(doc, "The running app does not show accuracy, sensitivity, specificity or confidence.")
    bullet(doc, "Direct float() conversion means that blank or non-numeric input can crash the page.")
    bullet(doc, "The result is a prediction message, not a clinical diagnosis workflow, laboratory confirmation or treatment plan.")
    bullet(doc, "No authentication, database, audit logging, rate limiting or monitoring is present in app.py.")
    bullet(doc, "Pickle artifacts should be treated as trusted files. Loading untrusted pickle data is a security risk in general and must not be done in production.")
    bullet(doc, "Parkinson’s training data is small (195 rows) and imbalanced (147 vs 48). Heart-disease test size is only 61 rows. These sample sizes limit claims of clinical reliability.")
    bullet(doc, "Several heart-disease fields are integer codes. The UI currently asks for them as free text, so a user can easily enter a value the model never saw.")
    add_text(
        doc,
        "Important academic disclaimer: this system is a student project and a prototype decision-support "
        "interface. It must not be used as a substitute for professional medical advice, diagnosis or "
        "treatment.",
        italic=True, before=8,
    )

    heading(doc, "12.  ASSUMPTIONS", 1)
    add_text(doc, "The implementation is based on the following assumptions:", after=4)
    numbered(doc, "The three CSV datasets are sufficiently structured and labelled for supervised binary classification.")
    numbered(doc, "The feature order used in the Streamlit forms matches the column order used to train each .sav file.")
    numbered(doc, "Historical patterns in the public datasets contain information that a linear SVM or logistic model can use.")
    numbered(doc, "An 80:20 split with a fixed random seed is an acceptable academic evaluation protocol for this demonstration.")
    numbered(doc, "Accuracy is an acceptable primary metric for the notebooks, even though class imbalance (especially in Parkinson’s) would justify extra metrics in a stronger study.")
    numbered(doc, "Users of the prototype can read the form labels and enter numeric values in the expected coding (for example 0/1 for sex).")
    numbered(doc, "The application is operated in a trusted environment where the .sav files have not been tampered with.")
    numbered(doc, "The system is used for learning, demonstration and prototype screening — not for unsupervised clinical decisions.")

    heading(doc, "13.  EXPECTED IMPACT AND BENEFITS", 1)
    add_text(
        doc,
        "In compact numbers the delivered prototype offers three disease modules, forty-three input "
        "fields across those modules, one unified interface, and no need to expose model source code "
        "to the end user.",
    )
    add_text(doc, "Benefits:", bold=True, after=4, align="left")
    bullet(doc, "Simplifies access to multiple ML predictors.")
    bullet(doc, "Standardises navigation and interaction.")
    bullet(doc, "Makes model inference accessible through a browser UI.")
    bullet(doc, "Supports academic demonstration of end-to-end ML deployment.")
    add_text(doc, "Potential impact:", bold=True, before=8, after=4, align="left")
    bullet(doc, "Useful as an educational / prototype decision-support application in a BCA laboratory or project viva.")
    bullet(doc, "Can be extended toward richer patient-facing workflows (history, PDF reports, range warnings).")
    bullet(doc, "Can become a foundation for monitoring, analytics and comparison of alternative models (Random Forest, XGBoost, calibrated probabilities).")

    heading(doc, "14.  FUTURE SCOPE AND FURTHER ENHANCEMENT", 1)
    add_text(
        doc,
        "The current implementation is a working prototype. Several enhancements can be introduced "
        "in later versions. They are grouped into five phases so that the roadmap remains practical.",
    )
    add_text(doc, "Phase 1 — Robust validation", bold=True, color=NAVY, after=4, align="left")
    add_text(
        doc,
        "Replace free text with number inputs and select boxes. Add required-field checks, medical "
        "range limits, and try/except around float conversion so that the user sees a clear error "
        "instead of a crash.",
    )
    add_text(doc, "Phase 2 — Better ML lifecycle", bold=True, color=NAVY, after=4, align="left")
    add_text(
        doc,
        "Add a documented training pipeline with preprocessing (imputation, scaling), additional "
        "metrics (precision, recall, F1, ROC-AUC, confusion matrix), class-imbalance handling, "
        "and versioned model files. Compare SVM and Logistic Regression with Random Forest, "
        "XGBoost or LightGBM. Address the Logistic Regression convergence warning by raising max_iter "
        "or scaling features.",
    )
    add_text(doc, "Phase 3 — User platform", bold=True, color=NAVY, after=4, align="left")
    add_text(
        doc,
        "Introduce authentication, optional patient profiles, prediction history and secure storage. "
        "This phase should not begin until privacy and consent rules are defined.",
    )
    add_text(doc, "Phase 4 — Clinical UX", bold=True, color=NAVY, after=4, align="left")
    add_text(
        doc,
        "Show prediction confidence, simple explanations (for example SHAP feature contributions), "
        "and a clear “not a diagnosis” banner. A clinician-review step would be required before any "
        "real-world health use.",
    )
    add_text(doc, "Phase 5 — Deployment", bold=True, color=NAVY, after=4, align="left")
    add_text(
        doc,
        "Containerise the application, add logging and monitoring, and serve models through a "
        "controlled API rather than loading pickle files in the web process. Production serving "
        "should prefer a safer serialisation format than pickle where possible.",
    )

    heading(doc, "15.  PROJECT REPOSITORY LOCATION", 1)
    add_text(
        doc,
        "The project artifacts for this academic submission are maintained as follows. The Streamlit "
        "application structure, training notebooks and datasets follow the well-known open-source "
        "Multiple Disease Prediction Streamlit App, which is cited in the references. The present "
        "report, analysis, diagrams and college documentation have been prepared for Teerthanker "
        "Mahaveer University submission.",
    )
    table = doc.add_table(rows=4, cols=5)
    heads = ["S#", "Project artifact (soft copy)", "Location", "Verified by Project Guide", "Verified by Lab In-Charge"]
    for i, h in enumerate(heads):
        cell_text(table.rows[0].cells[i], h, bold=True, size=8, align="center", fill=NAVY_HEX, color=WHITE)
    arts = [
        ["1", "Project Report (this document, final version)", "Device folder / student submission", GUIDE2, ""],
        ["2", "Functional code — app.py, requirements.txt, saved models, datasets, training notebooks", "Device folder / project workspace", GUIDE2, ""],
        ["3", "Reference source repository (application structure and notebooks)", "https://github.com/siddhardhan23/multiple-disease-prediction-streamlit-app", GUIDE2, ""],
    ]
    for r, row in enumerate(arts, start=1):
        fill = LIGHT_HEX if r % 2 else "FFFFFF"
        for c, val in enumerate(row):
            cell_text(table.rows[r].cells[c], val, size=8, fill=fill, align="center" if c != 1 and c != 2 else "left")
    style_table(table, header=False, col_widths=[0.5, 2.0, 1.9, 1.2, 1.0])
    caption(doc, "Table 10: Project artifacts and repository location (edit the location column if you upload the code to your own GitHub account).")
    add_text(
        doc,
        "If you later publish your own GitHub copy, replace row 2 with that URL and keep row 3 as "
        "the cited reference implementation.",
        size=11, italic=True, color=MUTED,
    )

    heading(doc, "16.  DEFINITIONS, ACRONYMS AND ABBREVIATIONS", 1)
    table = doc.add_table(rows=1, cols=2)
    cell_text(table.rows[0].cells[0], "Abbreviation", bold=True, size=11, align="center", fill=NAVY_HEX, color=WHITE)
    cell_text(table.rows[0].cells[1], "Description", bold=True, size=11, align="center", fill=NAVY_HEX, color=WHITE)
    abbr = [
        ("AI", "Artificial Intelligence"),
        ("ML", "Machine Learning"),
        ("SVM", "Support Vector Machine"),
        ("SVC", "Support Vector Classifier (scikit-learn class)"),
        ("LR", "Logistic Regression"),
        ("UI", "User Interface"),
        ("HLD", "High-Level Design"),
        ("LLD", "Low-Level Design"),
        ("DFD", "Data Flow Diagram"),
        ("CSV", "Comma-Separated Values"),
        ("SAV / PKL", "Serialized Python model file produced with pickle"),
        ("BMI", "Body Mass Index"),
        ("DPF", "Diabetes Pedigree Function"),
        ("MDVP", "Multi-Dimensional Voice Program (Parkinson’s acoustic features)"),
        ("NHR / HNR", "Noise-to-Harmonics Ratio / Harmonics-to-Noise Ratio"),
        ("RPDE", "Recurrence Period Density Entropy"),
        ("DFA", "Detrended Fluctuation Analysis"),
        ("PPE", "Pitch Period Entropy"),
        ("PIMA", "PIMA Indians Diabetes Database, the educational diabetes dataset used here"),
        ("UCI", "University of California, Irvine Machine Learning Repository"),
        ("API", "Application Programming Interface"),
        ("OCR", "Optical Character Recognition (future work only)"),
        ("SHAP", "SHapley Additive exPlanations (future explainability work)"),
        ("BCA", "Bachelor of Computer Application"),
        ("TMU", "Teerthanker Mahaveer University"),
    ]
    for i, (a, b) in enumerate(abbr, start=1):
        row = table.add_row().cells
        fill = LIGHT_HEX if i % 2 else "FFFFFF"
        cell_text(row[0], a, size=10, fill=fill, bold=True)
        cell_text(row[1], b, size=10, fill=fill)
    style_table(table, header=False, col_widths=[1.5, 5.0])
    caption(doc, "Table 11: Definitions, acronyms and abbreviations used in this report.")

    heading(doc, "17.  CONCLUSION", 1)
    add_text(
        doc,
        "The Multiple Disease Prediction System shows how three independent classification models "
        "can be delivered through one browser interface. The project starts from public tabular "
        "datasets, trains a linear SVM for diabetes, a logistic regression model for heart disease "
        "and a linear SVM for Parkinson’s, stores the fitted estimators as pickle artifacts, and "
        "serves them from a Streamlit application with sidebar navigation.",
    )
    add_text(
        doc,
        "The application-level design is consistent: every module converts form values to float, "
        "calls predict() on the matching .sav file, and returns a short classification message. "
        "That consistency is the main engineering contribution of the interface. The notebooks "
        "report test accuracies of approximately 77.27% (diabetes), 81.97% (heart disease) and "
        "87.18% (Parkinson’s). Those figures are academic baseline results on small public datasets; "
        "they are not a claim of clinical performance.",
    )
    add_text(
        doc,
        "The implementation with Python, scikit-learn and Streamlit is lightweight and suitable for "
        "laboratory demonstration and for a BCA project viva. Successful further use would require "
        "input validation, richer evaluation metrics, safer model serving, and a clear separation "
        "between a screening prototype and a medical device. Within its declared scope — a unified "
        "ML interface for three prediction tasks — the project meets its aim and provides a practical "
        "starting point for later enhancement.",
    )

    heading(doc, "18.  REFERENCES / BIBLIOGRAPHY", 1)
    refs = [
        "Siddhardhan, Multiple Disease Prediction Streamlit App, source code, training notebooks and datasets, GitHub repository, https://github.com/siddhardhan23/multiple-disease-prediction-streamlit-app",
        "Project source file analysed for this report: app.py — Multiple Disease Prediction System (Streamlit application, page configuration, sidebar, three prediction forms, pickle model loading and predict() calls).",
        "Project README.md — description of the Streamlit application, requirement installation, and the presence of training notebooks and datasets in their respective folders.",
        "Project requirements.txt — numpy==1.26.3, scikit-learn==1.3.2, streamlit==1.29.0, streamlit-option-menu==0.3.6.",
        "Streamlit documentation, application development and execution, https://docs.streamlit.io/",
        "Python documentation, Python programming language, https://docs.python.org/",
        "scikit-learn documentation, SVM, Logistic Regression, model evaluation and persistence, https://scikit-learn.org/",
        "NumPy documentation, numerical computing in Python, https://numpy.org/doc/",
        "Pandas documentation, data analysis and CSV handling, https://pandas.pydata.org/docs/",
        "Smith, J. W., Everhart, J. E., Dickson, W. C., Knowler, W. C. and Johannes, R. S. (1988). Using the ADAP learning algorithm to forecast the onset of diabetes mellitus. Proceedings of the Symposium on Computer Applications and Medical Care, 261–265. (PIMA diabetes dataset lineage.)",
        "UCI Machine Learning Repository, Heart Disease Data Set, https://archive.ics.uci.edu/dataset/45/heart+disease",
        "Little, M. A., McSharry, P. E., Hunter, E. J., Spielman, J. and Ramig, L. O. (2009). Suitability of dysphonia measurements for telemonitoring of Parkinson’s disease. IEEE Transactions on Biomedical Engineering, 56(4), 1015–1022. (Parkinson’s voice-measurement dataset lineage.)",
        "Cortes, C. and Vapnik, V. (1995). Support-vector networks. Machine Learning, 20, 273–297.",
        "Hosmer, D. W., Lemeshow, S. and Sturdivant, R. X. (2013). Applied Logistic Regression. Wiley.",
        "Pedregosa, F. et al. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825–2830.",
        "Pickle — Python object serialisation, https://docs.python.org/3/library/pickle.html",
        "World Health Organization resources on diabetes, cardiovascular disease and Parkinson’s disease (background motivation only; not used as a training source).",
    ]
    for i, ref in enumerate(refs, start=1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(f"[{i}]  {ref}")
        set_run_font(r, "Times New Roman", 11)

    add_text(
        doc,
        "Note: Accuracy figures quoted in Chapter 4 are those printed by the project training notebooks. "
        "They are not invented for this report. Where app.py itself is silent (for example on algorithm "
        "class names), the notebooks that produce diabetes_model.sav, heart_disease_model.sav and "
        "parkinsons_model.sav have been used as the source of truth.",
        size=11, italic=True, before=12, color=MUTED,
    )

    heading(doc, "APPENDIX A  —  FEATURE DICTIONARY AND RESULT MESSAGES", 1)
    add_text(
        doc,
        "This appendix restates the mapping from model output to on-screen text, as implemented in app.py.",
    )
    table = doc.add_table(rows=4, cols=3)
    for i, h in enumerate(["Module", "If predict() returns 1", "If predict() returns 0"]):
        cell_text(table.rows[0].cells[i], h, bold=True, size=10, align="center", fill=NAVY_HEX, color=WHITE)
    msgs = [
        ["Diabetes", "The person is diabetic", "The person is not diabetic"],
        ["Heart disease", "The person is having heart disease", "The person does not have any heart disease"],
        ["Parkinson’s", "The person has Parkinson's disease", "The person does not have Parkinson's disease"],
    ]
    for r, row in enumerate(msgs, start=1):
        fill = LIGHT_HEX if r % 2 else "FFFFFF"
        for c, val in enumerate(row):
            cell_text(table.rows[r].cells[c], val, size=10, fill=fill)
    style_table(table, header=False, col_widths=[1.6, 2.5, 2.5])
    caption(doc, "Appendix Table: Classification messages produced by app.py.")

    add_text(
        doc,
        "Sample predictive check used in the diabetes notebook (not a clinical case): input "
        "(5, 166, 72, 19, 175, 25.8, 0.587, 51) produced class 1 — “The person is diabetic”.",
    )
    add_text(
        doc,
        "Sample predictive check used in the heart notebook: input "
        "(62, 0, 0, 140, 268, 0, 0, 160, 0, 3.6, 0, 2, 2) produced class 0 — "
        "“The Person does not have a Heart Disease”.",
    )

    heading(doc, "APPENDIX B  —  SOFTWARE DEPENDENCIES AND RUNTIME COMMAND", 1)
    add_text(
        doc,
        "The Streamlit application dependencies, taken from requirements.txt, are:",
    )
    table = doc.add_table(rows=5, cols=3)
    for i, h in enumerate(["Package", "Pinned version", "Purpose"]):
        cell_text(table.rows[0].cells[i], h, bold=True, size=11, align="center", fill=NAVY_HEX, color=WHITE)
    deps = [
        ["numpy", "1.26.3", "Numerical arrays for feature vectors"],
        ["scikit-learn", "1.3.2", "Saved SVM / Logistic Regression estimators"],
        ["streamlit", "1.29.0", "Web application runtime and widgets"],
        ["streamlit-option-menu", "0.3.6", "Sidebar navigation control"],
    ]
    for r, row in enumerate(deps, start=1):
        fill = LIGHT_HEX if r % 2 else "FFFFFF"
        for c, val in enumerate(row):
            cell_text(table.rows[r].cells[c], val, size=11, fill=fill, align="center" if c < 2 else "left")
    style_table(table, header=False, col_widths=[2.1, 1.6, 2.8])
    caption(doc, "Table 12: Application dependencies from requirements.txt.")

    add_text(doc, "Typical local run (for laboratory demonstration):", bold=True, before=8, after=4, align="left")
    add_text(
        doc,
        "pip install -r requirements.txt\nstreamlit run app.py",
        font="Courier New", size=11, align="left", line=1.3, before=2, after=8,
    )
    add_text(
        doc,
        "The training notebooks require additional libraries commonly used with Jupyter (pandas, "
        "and the same scikit-learn / NumPy stack). The README notes that extra libraries may be "
        "needed to run those notebooks.",
    )

    # closing page
    page_break(doc)
    add_text(doc, "", size=12, before=80)
    p = add_text(doc, "MULTIPLE DISEASE PREDICTION SYSTEM", size=20, bold=True, align="center",
                 before=40, after=4, color=NAVY)
    add_text(doc, "From user input  →  ML inference  →  understandable prediction output",
             size=13, italic=True, align="center", before=4, after=16, color=TEAL)
    table = doc.add_table(rows=1, cols=3)
    cell_text(table.rows[0].cells[0], "DIABETES\n8 features  •  SVM", bold=True, size=12, align="center", fill=NAVY_HEX, color=WHITE)
    cell_text(table.rows[0].cells[1], "HEART DISEASE\n13 features  •  Logistic Regression", bold=True, size=12, align="center", fill=TEAL_HEX, color=WHITE)
    cell_text(table.rows[0].cells[2], "PARKINSON’S\n22 features  •  SVM", bold=True, size=12, align="center", fill=NAVY_HEX, color=WHITE)
    style_table(table, header=False, col_widths=[2.15, 2.15, 2.15])
    add_text(doc, "UNIFIED MACHINE-LEARNING INTERFACE", size=14, bold=True, align="center",
             before=20, after=20, color=NAVY)
    p = add_text(doc, "", after=4)
    add_hline(p, GOLD_HEX, "12")
    add_text(doc, "THANK YOU", size=22, bold=True, align="center", before=16, after=4, color=NAVY)
    add_text(doc, "Questions and Discussion", size=13, italic=True, align="center", after=12)
    add_text(doc, f"{STUDENT}  •  {ENROLL}  •  {PROGRAM}  •  Section {SECTION}",
             size=11, align="center", color=MUTED)
    add_text(doc, f"{FACULTY}\n{UNIV}", size=11, align="center", color=MUTED)
    add_text(doc, MONTH, size=11, align="center", color=MUTED)


def build():
    doc = Document()
    configure_styles(doc)
    doc.core_properties.author = STUDENT
    doc.core_properties.title = TITLE_LONG
    doc.core_properties.subject = "BCA Project Report — Multiple Disease Prediction System"
    doc.core_properties.category = "Academic Project Report"
    doc.core_properties.comments = f"Prepared for {UNIV}. Enrollment {ENROLL}."
    doc.core_properties.keywords = "Machine Learning, Streamlit, Diabetes, Heart Disease, Parkinson's, TMU, BCA"

    # --- Section 0: cover (no header, no page number look) ---
    sec0 = doc.sections[0]
    sec0.page_width = Inches(8.5)
    sec0.page_height = Inches(11)
    sec0.left_margin = Inches(1.0)
    sec0.right_margin = Inches(1.0)
    sec0.top_margin = Inches(0.9)
    sec0.bottom_margin = Inches(0.9)
    sec0.header_distance = Inches(0.4)
    sec0.footer_distance = Inches(0.4)
    sec0.different_first_page_header_footer = False
    setup_header_footer(sec0, show_header=False)
    set_page_number_format(sec0, "lowerRoman", 1)

    build_cover(doc)

    # --- front matter continues in same section with roman page numbers ---
    page_break(doc)
    setup_header_footer(sec0, show_header=True)
    # After cover we want headers. Cover already rendered without header text in first
    # paragraph of header - actually setup was called with show_header=False then we
    # need header from page 2. Using different first page is cleaner.

    # Rebuild header strategy: enable different first page so cover is clean.
    sec0.different_first_page_header_footer = True
    # first page header/footer empty
    fp_h = sec0.first_page_header
    fp_f = sec0.first_page_footer
    fp_h.is_linked_to_previous = False
    fp_f.is_linked_to_previous = False
    fp_h.paragraphs[0].clear()
    fp_f.paragraphs[0].clear()
    setup_header_footer(sec0, show_header=True)

    build_declaration(doc)
    page_break(doc)
    build_certificate(doc)
    page_break(doc)
    build_ack(doc)
    page_break(doc)
    build_abstract(doc)
    page_break(doc)
    build_toc(doc)
    page_break(doc)
    build_list_of_figures(doc)
    page_break(doc)
    build_list_of_tables(doc)

    # --- body section, arabic page numbers starting at 1 ---
    new_sec = doc.add_section()
    new_sec.page_width = Inches(8.5)
    new_sec.page_height = Inches(11)
    new_sec.left_margin = Inches(1.0)
    new_sec.right_margin = Inches(1.0)
    new_sec.top_margin = Inches(1.0)
    new_sec.bottom_margin = Inches(0.9)
    new_sec.header_distance = Inches(0.4)
    new_sec.footer_distance = Inches(0.4)
    new_sec.different_first_page_header_footer = False
    setup_header_footer(new_sec, show_header=True)
    set_page_number_format(new_sec, "decimal", 1)

    build_body(doc)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print("saved", OUT, "size", OUT.stat().st_size)


if __name__ == "__main__":
    build()
